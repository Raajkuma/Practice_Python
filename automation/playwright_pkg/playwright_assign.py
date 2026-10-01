"""
playwright_assign.py
WhatsApp Message Sender + Smart Data Extractor (attaches to YOUR existing Chrome)

Requirements:
    pip install -U playwright openpyxl

Chrome setup (one time, Chrome 144+):
    1. Open your normal Chrome (the one where WhatsApp Web is logged in).
    2. Go to: chrome://inspect/#remote-debugging
    3. Enable "Allow remote debugging for this browser instance".
    4. Keep Chrome open, then run this script.
       (Chrome may show an "Allow" prompt the first time - click Allow.)

Input  : contacts.xlsx  ->  columns: Name, Phone, Message (Message optional)
Output : reports/whatsapp_report_YYYY-MM-DD.json / .xlsx
         screenshots/YYYY-MM-DD/<index>_<name>.png

How a chat is opened (the fix):
    Every contact is opened with  https://web.whatsapp.com/send?phone=<digits>&text=<msg>
    This is a full page load, so the composer on screen ALWAYS belongs to the
    requested number. The old version searched the sidebar and then demanded that
    the chat header contain the Excel "Name" - which fails whenever the name saved
    in your phone differs from the Excel name, so nothing was ever sent.
"""

from __future__ import annotations

import json
import os
import random
import re
import time
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import quote

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from playwright.sync_api import (
    Locator,
    Page,
    TimeoutError as PlaywrightTimeoutError,
    sync_playwright,
)

# ============================================================
# CONFIGURATION
# ============================================================

# Everything is anchored to the folder that contains this script, so the
# reports/screenshots are always next to it no matter where you launch Python from.
BASE_DIR = Path(__file__).resolve().parent
CONTACTS_FILE = BASE_DIR / "contacts.xlsx"
REPORT_DIR = BASE_DIR / "reports"
SCREENSHOT_ROOT = BASE_DIR / "screenshots"

WHATSAPP_URL = "https://web.whatsapp.com/"
LOGIN_WAIT_SECONDS = 180
OPEN_CHAT_TIMEOUT_SECONDS = 60
DELIVERY_WAIT_SECONDS = 30
SEND_CONFIRM_TIMEOUT_SECONDS = 25

# How many of THEIR latest messages to read back into the report (assignment: 3).
LAST_N_MESSAGES = 3

# Added to 10-digit numbers that have no country code. "" disables this.
DEFAULT_COUNTRY_CODE = "91"

# Fallback if the chrome://inspect toggle file can't be found.
FALLBACK_CDP_URL = "http://127.0.0.1:9222"
# Set if your Chrome profile folder is in a non-standard place.
CHROME_USER_DATA_DIR = os.environ.get("CHROME_USER_DATA_DIR", "")

ACTION_DELAY_MIN = 2.0
ACTION_DELAY_MAX = 5.0
BATCH_PAUSE_EVERY = 10
BATCH_PAUSE_MIN = 10.0
BATCH_PAUSE_MAX = 20.0

DEFAULT_MESSAGE = (
    "Hello {name}, this is your daily update. "
    "Please let me know if you need anything."
)

# ============================================================
# SELECTORS (maintenance points - each list has fallbacks)
# ============================================================

MESSAGE_BOX_SELECTORS = [
    'footer div[contenteditable="true"][role="textbox"]',
    'footer div[contenteditable="true"]',
    'div[contenteditable="true"][data-tab="10"]',
    'div[contenteditable="true"][aria-label*="message" i]',
]
LOGGED_IN_SELECTORS = [
    "#pane-side",
    'div[contenteditable="true"][data-tab="3"]',
    'div[contenteditable="true"][aria-label*="Search" i]',
    'input[placeholder*="Search" i]',
]
QR_SELECTORS = [
    'canvas[aria-label*="QR" i]',
    'canvas[aria-label*="scan" i]',
    "div[data-ref] canvas",
]
SEND_BUTTON_SELECTORS = [
    'button[aria-label="Send"]',
    'button[aria-label*="Send" i]',
    'span[data-icon="send"]',
    '[data-testid="send"]',
]
# Clock icon shown on a message that has not left the browser yet.
PENDING_ICON = 'span[data-icon="msg-time"]'

NOT_FOUND_PHRASES = [
    "phone number shared via url is invalid",
    "phone number isn't on whatsapp",
    "not on whatsapp",
]


class ContactNotFoundError(Exception):
    """The number is invalid / not on WhatsApp."""


# ============================================================
# HELPERS
# ============================================================

def human_delay(lo: float = ACTION_DELAY_MIN, hi: float = ACTION_DELAY_MAX) -> None:
    time.sleep(random.uniform(lo, hi))


def sanitize_filename(value: str) -> str:
    value = re.sub(r'[<>:"/\\|?*\x00-\x1F]', "_", value)
    value = re.sub(r"\s+", "_", value).strip(" ._")
    return value[:100] or "contact"


def cell_to_text(value: Any) -> str:
    """Excel gives phone numbers as int/float; avoid '919876543210.0'."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_phone(phone: str) -> str:
    """Digits only, with country code added for bare 10-digit numbers."""
    raw = (phone or "").strip()
    has_plus = raw.startswith("+")
    digits = re.sub(r"\D", "", raw)
    if digits.startswith("00"):
        return digits[2:]
    if has_plus or not DEFAULT_COUNTRY_CODE:
        return digits
    if len(digits) == 10:
        return DEFAULT_COUNTRY_CODE + digits
    if len(digits) == 11 and digits.startswith("0"):
        return DEFAULT_COUNTRY_CODE + digits[1:]
    return digits


def norm_text(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip().lower()


def read_contacts(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        raise FileNotFoundError(f"Input file not found: {path.resolve()}")

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = wb.active.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration as exc:
            raise ValueError("contacts.xlsx is empty.") from exc

        headers = {cell_to_text(h).lower(): i for i, h in enumerate(raw_headers)}
        missing = {"name", "phone"} - set(headers)
        if missing:
            raise ValueError(
                "contacts.xlsx is missing required columns: "
                + ", ".join(sorted(missing))
            )

        def get(row, key):
            i = headers.get(key)
            return cell_to_text(row[i]) if i is not None and i < len(row) else ""

        contacts = []
        for row_number, row in enumerate(rows, start=2):
            name, phone, message = get(row, "name"), get(row, "phone"), get(row, "message")
            if not (name or phone or message):
                continue
            if not name:
                raise ValueError(f"Row {row_number}: Name is required.")
            if not phone:
                raise ValueError(f"Row {row_number}: Phone is required.")
            contacts.append({"Name": name, "Phone": phone, "Message": message})

        if not contacts:
            raise ValueError("No contact rows found in contacts.xlsx.")
        return contacts
    finally:
        wb.close()


def wait_visible(page: Page, selectors: list[str], timeout_ms: int) -> Locator | None:
    loc = page.locator(", ".join(selectors)).first
    try:
        loc.wait_for(state="visible", timeout=timeout_ms)
        return loc
    except PlaywrightTimeoutError:
        return None


def any_visible(page: Page, selectors: list[str]) -> bool:
    try:
        return page.locator(", ".join(selectors)).first.is_visible()
    except Exception:
        return False


def body_contains_not_found(page: Page) -> bool:
    try:
        text = page.locator("body").inner_text(timeout=2000).lower()
    except Exception:
        return False
    return any(p in text for p in NOT_FOUND_PHRASES)








# ============================================================
# CONNECT TO YOUR EXISTING CHROME
# ============================================================

def _user_data_dirs() -> list[Path]:
    dirs: list[Path] = []
    if CHROME_USER_DATA_DIR:
        dirs.append(Path(CHROME_USER_DATA_DIR))
    local = os.environ.get("LOCALAPPDATA")
    if local:
        dirs.append(Path(local) / "Google" / "Chrome" / "User Data")
    home = Path.home()
    dirs.append(home / "Library" / "Application Support" / "Google" / "Chrome")
    dirs.append(home / ".config" / "google-chrome")
    return dirs


def _devtools_ws_endpoints() -> list[str]:
    """
    Chrome 144+ (chrome://inspect toggle) writes DevToolsActivePort:
        line 1 = port, line 2 = websocket path.
    """
    endpoints = []
    for d in _user_data_dirs():
        f = d / "DevToolsActivePort"
        try:
            if f.exists():
                lines = f.read_text(encoding="utf-8").splitlines()
                if len(lines) >= 2 and lines[0].strip().isdigit():
                    endpoints.append(f"ws://127.0.0.1:{lines[0].strip()}{lines[1].strip()}")
        except OSError:
            continue
    return endpoints


def connect_existing_chrome(playwright):
    errors = []
    candidates = _devtools_ws_endpoints() + [FALLBACK_CDP_URL]
    for endpoint in candidates:
        try:
            print(f"Connecting to Chrome: {endpoint}")
            return playwright.chromium.connect_over_cdp(endpoint, timeout=30000)
        except Exception as exc:  # noqa: BLE001
            errors.append(f"{endpoint}: {exc}")

    raise RuntimeError(
        "Could not connect to your existing Chrome.\n"
        "  1. Open normal Chrome and keep it open.\n"
        "  2. Go to chrome://inspect/#remote-debugging\n"
        "  3. Enable 'Allow remote debugging for this browser instance'.\n"
        "  4. Click 'Allow' if Chrome shows a prompt, then run again.\n"
        "  (If your profile is in a custom folder, set CHROME_USER_DATA_DIR.)\n\n"
        "Attempts:\n  " + "\n  ".join(errors)
    )


def get_work_page(browser) -> Page:
    if not browser.contexts:
        raise RuntimeError("Connected Chrome has no browser context.")
    context = browser.contexts[0]

    wa_pages = [p for p in context.pages if "web.whatsapp.com" in (p.url or "")]
    # Reuse a WhatsApp tab, otherwise open a NEW tab (never hijack another tab).
    page = wa_pages[0] if wa_pages else context.new_page()

    page.set_default_timeout(15000)
    # WhatsApp may show a "Leave site?" dialog when we navigate - accept it.
    page.on("dialog", lambda d: d.accept())
    page.bring_to_front()
    return page


def ensure_whatsapp_login(page: Page) -> None:
    print("\nOpening WhatsApp Web...")
    if "web.whatsapp.com" not in (page.url or ""):
        page.goto(WHATSAPP_URL, wait_until="domcontentloaded")

    deadline = time.monotonic() + LOGIN_WAIT_SECONDS
    announced = False
    while time.monotonic() < deadline:
        if any_visible(page, LOGGED_IN_SELECTORS) or any_visible(page, MESSAGE_BOX_SELECTORS):
            print("WhatsApp Web session is ready.")
            return
        if not announced and any_visible(page, QR_SELECTORS):
            print(f"QR code shown - scan it with your phone (waiting up to {LOGIN_WAIT_SECONDS}s)...")
            announced = True
        page.wait_for_timeout(1000)

    if any_visible(page, QR_SELECTORS):
        raise RuntimeError("WhatsApp Web is not logged in (QR code still showing).")
    print("Warning: could not positively detect the chat list; continuing anyway.")


# ============================================================
# OPEN CHAT / SEND / CONFIRM
# ============================================================

def open_chat(page: Page, digits: str, message: str, name: str, phone: str) -> Locator:
    """Load the chat for this number (message prefilled) and return the composer."""
    if not digits:
        raise ContactNotFoundError(f"Invalid phone number for {name}: {phone!r}")

    url = f"{WHATSAPP_URL}send?phone={digits}&text={quote(message, safe='')}"
    page.goto(url, wait_until="domcontentloaded")

    composer_css = ", ".join(MESSAGE_BOX_SELECTORS)
    deadline = time.monotonic() + OPEN_CHAT_TIMEOUT_SECONDS
    while time.monotonic() < deadline:
        if body_contains_not_found(page):
            raise ContactNotFoundError(f"Number not on WhatsApp / invalid: {name} ({phone})")
        try:
            composer = page.locator(composer_css).first
            if composer.is_visible():
                return composer
        except Exception:
            pass
        page.wait_for_timeout(500)

    raise PlaywrightTimeoutError(f"Chat did not open within {OPEN_CHAT_TIMEOUT_SECONDS}s for {name} ({phone})")


def _composer_text(composer: Locator) -> str:
    try:
        return norm_text(composer.inner_text(timeout=1500))
    except Exception:
        return ""


def _composer_ok(actual: str, expected: str) -> bool:
    if not actual:
        return False
    return actual == expected or expected[:20] in actual or actual[:20] in expected


def ensure_composer_text(page: Page, composer: Locator, message: str) -> None:
    """Use the URL-prefilled text; if missing, type it ourselves."""
    expected = norm_text(message)

    for _ in range(10):  # up to ~5 s for the prefill to appear
        if _composer_ok(_composer_text(composer), expected):
            return
        page.wait_for_timeout(500)

    composer.click()
    page.keyboard.press("ControlOrMeta+A")
    page.keyboard.press("Delete")
    lines = message.split("\n")
    for i, line in enumerate(lines):
        if line:
            page.keyboard.insert_text(line)
        if i < len(lines) - 1:
            page.keyboard.press("Shift+Enter")  # newline without sending
    page.wait_for_timeout(500)

    if not _composer_ok(_composer_text(composer), expected):
        raise RuntimeError("Could not place the message into the WhatsApp composer.")









def _composer_holds_message(composer: Locator, expected: str) -> bool:
    """True while the composer still contains our (unsent) text.
    A placeholder such as 'Type a message' does NOT match, so it counts as empty."""
    return _composer_ok(_composer_text(composer), expected)


def message_visible_in_chat(page: Page, expected: str) -> bool:
    """Last-resort check: is our text visible in the conversation (outside the composer)?"""
    probe = expected[:40]
    if not probe:
        return False
    try:
        return bool(page.evaluate(
            """(probe) => {
                const m = (document.querySelector('#main') || document.body).cloneNode(true);
                m.querySelectorAll('footer').forEach(f => f.remove());
                return (m.textContent || '').replace(/\\s+/g, ' ').toLowerCase().includes(probe);
            }""",
            probe,
        ))
    except Exception:
        return False


def dump_debug(page: Page, path: Path) -> None:
    """Save the chat DOM so selectors can be fixed if WhatsApp changed."""
    try:
        html = page.evaluate("() => (document.querySelector('#main') || document.body).outerHTML")
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(html[:400000], encoding="utf-8")
        print(f"  Debug HTML saved: {path}")
    except Exception:
        pass


def _click_send_button(page: Page) -> bool:
    for sel in SEND_BUTTON_SELECTORS:
        try:
            btn = page.locator(sel).first
            if btn.count() > 0 and btn.is_visible():
                btn.click()
                return True
        except Exception:
            continue
    return False


def wait_until_delivered(page: Page) -> str:
    """Stay on the chat until no message shows the pending clock icon."""
    page.wait_for_timeout(1500)
    end = time.monotonic() + DELIVERY_WAIT_SECONDS
    while time.monotonic() < end:
        try:
            if page.locator(PENDING_ICON).count() == 0:
                return ""
        except Exception:
            return ""
        page.wait_for_timeout(500)
    return "Message still pending (clock icon) after wait - check your phone's internet."


def send_and_confirm(page: Page, composer: Locator, message: str, debug_path: Path | None = None) -> str:
    """
    Press Enter and treat the send as successful as soon as the composer no longer
    holds our text. (It held the text just before, so if it is gone, WhatsApp took it.)
    Returns a warning string ('' if clean). Raises only if the text never left the composer.
    """
    expected = norm_text(message)

    human_delay(1.0, 2.5)
    composer.click()
    composer.press("Enter")

    started = time.monotonic()
    deadline = started + SEND_CONFIRM_TIMEOUT_SECONDS
    clicked_button = False
    confirmed = False
    while time.monotonic() < deadline:
        if not _composer_holds_message(composer, expected):
            confirmed = True
            break
        if not clicked_button and time.monotonic() - started > 3:
            clicked_button = _click_send_button(page)
        page.wait_for_timeout(400)

    if not confirmed and message_visible_in_chat(page, expected):
        confirmed = True

    if not confirmed:
        if debug_path:
            dump_debug(page, debug_path)
        raise PlaywrightTimeoutError("WhatsApp did not accept the message (it is still in the composer).")

    return wait_until_delivered(page)


def take_screenshot(page: Page, path: Path) -> bool:
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        page.wait_for_timeout(500)
        page.screenshot(path=str(path))
        return True
    except Exception:
        return False


# ============================================================
# SMART DATA EXTRACTION
# ============================================================


EXTRACT_JS = r"""
() => {
  const main = document.querySelector('#main') || document.querySelector('div[role="application"]') || document.body;
  const clean = (node) => {
    const c = node.cloneNode(true);
    c.querySelectorAll('img[alt]').forEach(i => i.replaceWith(document.createTextNode(i.getAttribute('alt') || '')));
    return (c.textContent || '').replace(/\s+/g, ' ').trim();
  };
  const textOf = (el) => {
    const pre = el.querySelector('[data-pre-plain-text]');
    if (pre) {
      const sel = pre.querySelector('span.selectable-text');
      return { text: clean(sel || pre), meta: (pre.getAttribute('data-pre-plain-text') || '').trim() };
    }
    const s = el.querySelector('span.selectable-text');
    if (s) return { text: clean(s), meta: '' };
    // media / sticker / call etc.: use whatever text is there minus the trailing clock time
    let t = clean(el).replace(/\s*\d{1,2}:\d{2}\s*([AaPp][Mm])?\s*$/, '').trim();
    return { text: t || '[non-text message]', meta: '' };
  };
  const outermost = (els) => els.filter(e => !els.some(o => o !== e && o.contains(e)));

  let mode = 'data-id';
  let els = [...main.querySelectorAll('div[data-id]')].filter(e => /^(true|false)_/.test(e.getAttribute('data-id')));
  let dirOf = (e) => e.getAttribute('data-id').startsWith('true_') ? 'out' : 'in';

  if (!els.length) {
    mode = 'class';
    els = [...main.querySelectorAll('.message-in, .message-out')];
    dirOf = (e) => e.classList.contains('message-out') ? 'out' : 'in';
  }
  if (!els.length) {
    mode = 'position';   // last resort: left-aligned bubbles = incoming, right-aligned = outgoing
    const rows = [...main.querySelectorAll('[data-pre-plain-text]')]
      .map(n => n.closest('div[role="row"]') || n.parentElement.parentElement);
    els = [...new Set(rows)].filter(Boolean);
    const mr = main.getBoundingClientRect();
    dirOf = (e) => {
      const b = e.querySelector('[data-pre-plain-text]') || e;
      const r = b.getBoundingClientRect();
      return ((r.left + r.right) / 2) > ((mr.left + mr.right) / 2) ? 'out' : 'in';
    };
  }
  els = outermost(els);
  const items = els.map(e => Object.assign({ dir: dirOf(e) }, textOf(e)));
  return { mode, total: items.length, items };
}
"""


SCROLL_JS = r"""
(mode) => {
  const main = document.querySelector('#main') || document.body;
  const first = main.querySelector('[data-pre-plain-text], div[data-id], div[role="row"], .message-in, .message-out');
  let el = first ? first.parentElement : null;
  while (el && el !== document.body) {
    const oy = getComputedStyle(el).overflowY;
    if (el.scrollHeight > el.clientHeight + 5 && (oy === 'auto' || oy === 'scroll')) break;
    el = el.parentElement;
  }
  if (!el || el === document.body) return -1;
  if (mode === 'up') el.scrollTop = 0; else el.scrollTop = el.scrollHeight;
  return el.scrollHeight;
}
"""


def _format_incoming(item: dict[str, Any]) -> str:
    text = (item.get("text") or "").strip()
    m = re.match(r"\[(.*?)\]", item.get("meta") or "")
    return f"[{m.group(1)}] {text}" if m else text


def _scan_chat(page: Page) -> dict[str, Any]:
    try:
        return page.evaluate(EXTRACT_JS)
    except Exception:
        return {"total": 0, "items": [], "mode": ""}


def _incoming_only(data: dict[str, Any]) -> list[dict[str, Any]]:
    return [i for i in data["items"] if i["dir"] == "in" and (i.get("text") or "").strip()]


def extract_last_messages_from_them(page: Page, n: int = LAST_N_MESSAGES,
                                    debug_path: Path | None = None) -> tuple[list[str], str]:
    """
    Read the last n INCOMING messages of the chat that is currently open.
    - waits for the history to finish loading
    - scrolls up (a few times) if fewer than n incoming messages are loaded
    Returns (messages, note). Raises only if the chat could not be read at all.
    """
    data: dict[str, Any] = {"total": 0, "items": [], "mode": ""}
    last_total = -1
    for _ in range(10):                      # wait for history to load
        page.wait_for_timeout(800)
        data = _scan_chat(page)
        if data["total"] > 0 and data["total"] == last_total:
            break
        last_total = data["total"]

    if data["total"] == 0:
        if debug_path:
            dump_debug(page, debug_path)
        raise RuntimeError("No messages could be read from the chat "
                           "(selectors may be outdated; see the saved debug HTML).")

    # Not enough incoming messages on screen -> load older history.
    for _ in range(6):
        if len(_incoming_only(data)) >= n:
            break
        try:
            if page.evaluate(SCROLL_JS, "up") == -1:
                break
        except Exception:
            break
        grew = False
        for _ in range(5):                   # older messages load asynchronously: be patient
            page.wait_for_timeout(800)
            candidate = _scan_chat(page)
            if candidate["total"] > data["total"]:
                data, grew = candidate, True
                page.wait_for_timeout(600)   # let the rest of the batch land
                data = _scan_chat(page)
                break
        if not grew:
            break
    try:
        page.evaluate(SCROLL_JS, "down")
    except Exception:
        pass

    incoming = _incoming_only(data)
    print(f"  Chat scan: mode={data['mode']}, messages seen={data['total']}, from them={len(incoming)}")
    messages = [_format_incoming(i) for i in incoming[-n:]]
    note = "" if messages else "Chat opened, but this contact has not sent any messages (nothing to extract)."
    return messages, note


def reopen_chat_for_reading(page: Page, digits: str, name: str, phone: str) -> None:
    """Open the contact's chat fresh (no text prefilled) so we read the real history."""
    page.goto(f"{WHATSAPP_URL}send?phone={digits}", wait_until="domcontentloaded")
    composer_css = ", ".join(MESSAGE_BOX_SELECTORS)
    deadline = time.monotonic() + OPEN_CHAT_TIMEOUT_SECONDS
    while time.monotonic() < deadline:
        if body_contains_not_found(page):
            raise ContactNotFoundError(f"Number not on WhatsApp: {name} ({phone})")
        try:
            if page.locator(composer_css).first.is_visible():
                return
        except Exception:
            pass
        page.wait_for_timeout(500)
    raise PlaywrightTimeoutError("Could not reopen the chat to read messages.")


# ============================================================
# REPORTS
# ============================================================

def build_result_row(*, index: int, name: str, phone: str, template: str) -> dict[str, Any]:
    return {
        "index": index,
        "name": name,
        "phone": phone,
        "dialed_number": "",
        "message_template": template,
        "personalized_message": "",
        "status": "Pending",
        "sent_at": "",
        "screenshot": "",
        "last_3_messages": [],
        "extraction_note": "",
        "warning": "",
        "extraction_error": "",
        "error": "",
    }


def save_json_report(path: Path, report: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")


def write_with_fallback(writer, path: Path, *args) -> Path | None:
    """
    Run writer(path, *args). If the file is locked (e.g. the .xlsx is open in Excel)
    write to a timestamped sibling instead, so a report is NEVER lost or crashes the run.
    """
    try:
        writer(path, *args)
        return path
    except Exception as first:  # noqa: BLE001
        alt = path.with_name(f"{path.stem}_{datetime.now():%H%M%S}{path.suffix}")
        try:
            writer(alt, *args)
            print(f"  WARNING: could not write {path.name} ({first}); wrote {alt.name} instead.")
            return alt
        except Exception as second:  # noqa: BLE001
            print(f"  ERROR: could not write report {path}: {second}")
            return None


def style_header(ws) -> None:
    fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def save_excel_report(path: Path, results: list[dict[str, Any]], meta: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Name", "Phone", "Status", "Message Sent", "Sent At", "Screenshot",
               "Last 3 Messages (from them)", "Warning", "Extraction Note", "Extraction Error", "Error"])
    for r in results:
        ws.append([r["name"], r["phone"], r["status"], r["personalized_message"], r["sent_at"],
                   r["screenshot"], "\n".join(r["last_3_messages"]), r["warning"],
                   r["extraction_note"], r["extraction_error"], r["error"]])
    style_header(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate([22, 18, 14, 55, 24, 55, 60, 40, 40, 40, 55], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)

    m = wb.create_sheet("Run Info")
    m.append(["Field", "Value"])
    style_header(m)
    for k, v in meta.items():
        m.append([k, v])
    m.column_dimensions["A"].width = 28
    m.column_dimensions["B"].width = 60
    wb.save(path)


# ============================================================
# MAIN
# ============================================================

def run() -> None:
    run_date = datetime.now().strftime("%Y-%m-%d")
    started = datetime.now().isoformat(timespec="seconds")
    screenshot_dir = SCREENSHOT_ROOT / run_date
    debug_dir = screenshot_dir / "debug"
    json_path = REPORT_DIR / f"whatsapp_report_{run_date}.json"
    xlsx_path = REPORT_DIR / f"whatsapp_report_{run_date}.xlsx"

    results: list[dict[str, Any]] = []
    contacts: list[dict[str, str]] = []
    fatal: list[str] = []
    final_paths: dict[str, Path | None] = {"json": json_path, "xlsx": xlsx_path}

    def write_reports() -> None:
        counts = {s: sum(1 for r in results if r["status"] == s)
                  for s in ("Sent", "Not Found", "Failed")}
        meta = {
            "run_date": run_date,
            "run_started_at": started,
            "total_contacts": len(contacts),
            "processed": len(results),
            "sent": counts["Sent"],
            "not_found": counts["Not Found"],
            "failed": counts["Failed"],
            "messages_extracted_for": sum(1 for r in results if r["last_3_messages"]),
            "fatal_error": "; ".join(fatal),
            "input_file": str(CONTACTS_FILE),
            "json_report": str(json_path),
            "excel_report": str(xlsx_path),
        }
        # If a file is locked we switch to a fallback name and keep using THAT name
        # for the rest of the run (instead of creating a new file every time).
        final_paths["json"] = write_with_fallback(
            save_json_report, final_paths["json"] or json_path,
            {"report": meta, "results": results}) or final_paths["json"]
        final_paths["xlsx"] = write_with_fallback(
            save_excel_report, final_paths["xlsx"] or xlsx_path, results, meta) or final_paths["xlsx"]

    try:
        contacts = read_contacts(CONTACTS_FILE)
        print(f"\nLoaded {len(contacts)} contacts from {CONTACTS_FILE}")
        print(f"Reports will be saved in: {REPORT_DIR}")

        with sync_playwright() as pw:
            browser = connect_existing_chrome(pw)
            page = get_work_page(browser)
            ensure_whatsapp_login(page)

            for index, contact in enumerate(contacts, start=1):
                name, phone = contact["Name"], contact["Phone"]
                template = contact.get("Message", "")
                digits = normalize_phone(phone)
                tag = f"{index:03d}_{sanitize_filename(name)}"

                result = build_result_row(index=index, name=name, phone=phone, template=template)
                message = (template or DEFAULT_MESSAGE).replace("{name}", name)
                result["personalized_message"] = message
                result["dialed_number"] = digits

                print(f"\n[{index}/{len(contacts)}] {name} ({phone} -> {digits})")
                chat_opened = False

                # ---------- 1) SEND ----------
                try:
                    human_delay()
                    composer = open_chat(page, digits, message, name, phone)
                    chat_opened = True
                    ensure_composer_text(page, composer, message)
                    warning = send_and_confirm(page, composer, message, debug_dir / f"{tag}_send.html")

                    result["status"] = "Sent"
                    result["sent_at"] = datetime.now().isoformat(timespec="seconds")
                    result["warning"] = warning
                    print("  Message sent and confirmed." + (f"  WARNING: {warning}" if warning else ""))

                    shot = screenshot_dir / f"{tag}.png"
                    if take_screenshot(page, shot):
                        result["screenshot"] = str(shot)

                except ContactNotFoundError as exc:
                    result["status"], result["error"] = "Not Found", str(exc)
                    print(f"  Not found: {exc}")
                except PlaywrightTimeoutError as exc:
                    result["status"], result["error"] = "Failed", f"Timeout: {exc}"
                    print(f"  Timeout: {exc}")
                except Exception as exc:  # noqa: BLE001
                    result["status"], result["error"] = "Failed", f"{type(exc).__name__}: {exc}"
                    print(f"  Failed: {exc}")

                # ---------- 2) SMART EXTRACTION (own try: never changes the send status) ----------
                if chat_opened:
                    try:
                        print("  Reopening the contact to read their last messages...")
                        try:
                            reopen_chat_for_reading(page, digits, name, phone)
                        except Exception as exc:  # fall back to the chat already on screen
                            print(f"  (reopen failed: {exc}; reading the current view)")
                        msgs, note = extract_last_messages_from_them(
                            page, LAST_N_MESSAGES, debug_dir / f"{tag}_read.html")
                        result["last_3_messages"] = msgs
                        result["extraction_note"] = note
                        if msgs:
                            print(f"  Last {len(msgs)} message(s) from {name}:")
                            for m in msgs:
                                print(f"     - {m}")
                        else:
                            print(f"  {note}")
                    except Exception as exc:  # noqa: BLE001
                        result["extraction_error"] = f"{type(exc).__name__}: {exc}"
                        print(f"  Extraction warning: {exc}")

                results.append(result)
                write_reports()  # save JSON + Excel after EVERY contact

                if index % BATCH_PAUSE_EVERY == 0 and index < len(contacts):
                    pause = random.uniform(BATCH_PAUSE_MIN, BATCH_PAUSE_MAX)
                    print(f"  Batch pause: {pause:.1f}s")
                    time.sleep(pause)

            # Don't close the browser: it's your own Chrome.
            print("\nAutomation completed.")

    except BaseException as exc:  # noqa: BLE001  (also Ctrl+C: still save what we have)
        fatal.append(f"{type(exc).__name__}: {exc}")
        print(f"\nRun stopped early: {type(exc).__name__}: {exc}")
        raise
    finally:
        write_reports()  # ALWAYS produce both reports, even after an early stop
        print(f"\nJSON report : {final_paths['json']}")
        print(f"Excel report: {final_paths['xlsx']}")
        print("Summary -> " + ", ".join(
            f"{s}: {sum(1 for r in results if r['status'] == s)}"
            for s in ("Sent", "Not Found", "Failed")))


if __name__ == "__main__":
    run()